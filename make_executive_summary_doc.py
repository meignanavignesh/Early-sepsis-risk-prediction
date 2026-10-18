from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

output = "Early_Sepsis_Executive_Summary.docx"
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.85)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(8)
styles["Normal"].paragraph_format.line_spacing = 1.08

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(16)
run = title.add_run("Executive Summary")
run.bold = True
run.font.name = "Calibri"
run.font.size = Pt(16)

paragraphs = [
    "The project entitled “Early Sepsis Risk Prediction using Electronic Health Records and Explainable AI” focuses on the early identification of sepsis risk among intensive care unit patients using hourly electronic health record observations. The dataset contains approximately 40,336 patients and more than 1.5 million hourly observations, including vital signs, laboratory measurements, demographic information, ICU stay duration, and sepsis labels.",
    "The system performs data preprocessing, missing-value handling, and temporal feature engineering. Lag values, hourly changes, and 3-hour and 6-hour rolling statistics are generated to capture changes in a patient’s condition over time. An optimized XGBoost model uses these engineered features to estimate sepsis risk for each hourly observation.",
    "A decision layer classifies patients into different risk categories and generates alerts when the predicted probability exceeds the configured threshold. The Streamlit dashboard provides patient selection, hourly risk trends, patient summaries, alert information, and SHAP-based explanations of the features contributing to each prediction.",
    "The proposed system serves as a clinical decision-support prototype that can help healthcare professionals monitor patient deterioration and prioritize high-risk cases. Further clinical validation is required before real-world deployment.",
]

for text in paragraphs:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.core_properties.title = "Executive Summary - Early Sepsis Risk Prediction"
doc.core_properties.subject = "Project executive summary"
doc.core_properties.author = ""
doc.save(output)
print(output)
